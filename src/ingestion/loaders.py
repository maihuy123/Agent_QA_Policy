"""Đọc file PDF, DOCX, Markdown thành các phần văn bản kèm metadata nguồn."""

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import yaml

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown"}


@dataclass
class Section:
    """Một phần của tài liệu: một trang PDF hoặc một mục Markdown/DOCX."""

    text: str
    page: int = 0  # 0 nghĩa là định dạng không có khái niệm trang
    heading: str = ""


@dataclass
class LoadedDocument:
    """Tài liệu đã đọc: metadata cấp file và danh sách các phần."""

    source: str
    title: str
    author: str
    updated_at: str
    file_type: str
    sections: list[Section] = field(default_factory=list)


def load_document(path: str | Path) -> LoadedDocument:
    """Đọc một file theo phần mở rộng; báo lỗi nếu định dạng chưa hỗ trợ."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _load_pdf(path)
    if ext == ".docx":
        return _load_docx(path)
    if ext in (".md", ".markdown"):
        return _load_markdown(path)
    raise ValueError(f"Định dạng chưa hỗ trợ: {path.name}")


def _file_mtime(path: Path) -> str:
    return datetime.fromtimestamp(path.stat().st_mtime).date().isoformat()


def _format_date(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip() if value else ""


def _load_pdf(path: Path) -> LoadedDocument:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    meta = reader.metadata
    updated = ""
    if meta is not None:
        try:
            updated = _format_date(meta.modification_date or meta.creation_date)
        except Exception:
            # Ngày trong metadata PDF đôi khi sai định dạng, khi đó dùng ngày sửa file
            updated = ""
    sections = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            sections.append(Section(text=text, page=i))
    return LoadedDocument(
        source=path.name,
        title=(meta.title if meta and meta.title else path.stem),
        author=(meta.author if meta and meta.author else ""),
        updated_at=updated or _file_mtime(path),
        file_type="pdf",
        sections=sections,
    )


def _load_docx(path: Path) -> LoadedDocument:
    import docx
    from docx.table import Table

    document = docx.Document(str(path))
    props = document.core_properties
    sections: list[Section] = []
    heading, lines = "", []

    def flush():
        text = "\n".join(lines).strip()
        if text:
            sections.append(Section(text=text, heading=heading))

    # Duyệt đoạn văn và bảng theo đúng thứ tự trong file; tách mục theo style Heading
    # để bảng nằm dưới đúng tên mục chứa nó khi trích dẫn
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            lines.extend(_table_lines(block))
            continue
        text = block.text.strip()
        if not text:
            continue
        style = block.style.name.lower() if block.style is not None else ""
        if style.startswith("heading") or style == "title":
            flush()
            heading, lines = text, []
        else:
            lines.append(text)
    flush()

    return LoadedDocument(
        source=path.name,
        title=props.title or path.stem,
        author=props.author or "",
        updated_at=_format_date(props.modified) or _file_mtime(path),
        file_type="docx",
        sections=sections,
    )


def _table_lines(table) -> list[str]:
    """Đổi bảng thành các dòng "Tên cột: giá trị; ..." (dòng đầu là tiêu đề cột) để embedding hiểu được."""
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    rows = [r for r in rows if any(r)]
    if len(rows) < 2:
        return [" | ".join(r) for r in rows]
    header = rows[0]
    return ["; ".join(f"{h}: {v}" if h else v for h, v in zip(header, r) if v) for r in rows[1:]]


_FRONT_MATTER =re.compile(r"\A---\s*\n(.*?)\n---\s*\n", re.DOTALL)
_HEADING = re.compile(r"^(#{1,3})\s+(.*)$")


def _load_markdown(path: Path) -> LoadedDocument:
    raw = path.read_text(encoding="utf-8")
    meta: dict = {}
    match = _FRONT_MATTER.match(raw)
    if match:
        meta = yaml.safe_load(match.group(1)) or {}
        raw = raw[match.end() :]

    title = str(meta.get("title", "")).strip()
    sections: list[Section] = []
    heading, lines = "", []

    def flush():
        text = "\n".join(lines).strip()
        if text:
            sections.append(Section(text=text, heading=heading))

    for line in raw.splitlines():
        m = _HEADING.match(line)
        if m:
            flush()
            heading, lines = m.group(2).strip(), []
            # Tiêu đề cấp 1 đầu tiên được dùng làm tên tài liệu nếu front matter không có
            if m.group(1) == "#" and not title:
                title = heading
        else:
            lines.append(line)
    flush()

    return LoadedDocument(
        source=path.name,
        title=title or path.stem,
        author=str(meta.get("author", "")).strip(),
        updated_at=_format_date(meta.get("updated")) or _file_mtime(path),
        file_type="markdown",
        sections=sections,
    )
