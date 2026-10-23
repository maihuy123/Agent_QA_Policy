from datetime import datetime

import docx
import pytest

from src.ingestion.loaders import load_document


def test_markdown_front_matter_and_sections(tmp_path):
    path = tmp_path / "policy.md"
    path.write_text(
        "---\ntitle: Chính sách\nauthor: Phòng Nhân sự\nupdated: 2026-03-01\n---\n"
        "# Chính sách\nGiới thiệu.\n\n## Mục A\nNội dung A.\n\n## Mục B\nNội dung B.\n",
        encoding="utf-8",
    )
    doc = load_document(path)
    assert (doc.source, doc.title, doc.author, doc.updated_at) == ("policy.md", "Chính sách", "Phòng Nhân sự", "2026-03-01")
    assert [s.heading for s in doc.sections] == ["Chính sách", "Mục A", "Mục B"]
    assert doc.sections[1].text == "Nội dung A."


def test_markdown_without_front_matter_uses_h1_as_title(tmp_path):
    path = tmp_path / "note.md"
    path.write_text("# Ghi chú\nNội dung.", encoding="utf-8")
    doc = load_document(path)
    assert doc.title == "Ghi chú"
    assert doc.updated_at  # lấy theo ngày sửa file


def test_docx_headings_tables_and_properties(tmp_path):
    path = tmp_path / "guide.docx"
    d = docx.Document()
    d.add_heading("Hướng dẫn", level=0)
    d.add_heading("Bước 1", level=1)
    d.add_paragraph("Làm việc A.")
    table = d.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Đầu mối", "Liên hệ"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "IT", "máy lẻ 200"
    d.core_properties.author = "Phòng IT"
    d.core_properties.modified = datetime(2026, 4, 1)
    d.save(path)

    doc = load_document(path)
    assert doc.author == "Phòng IT"
    assert doc.updated_at == "2026-04-01"
    assert any(s.heading == "Bước 1" and "Làm việc A." in s.text for s in doc.sections)
    # Bảng nằm ngay sau đoạn của "Bước 1" nên thuộc mục đó
    assert any(s.heading == "Bước 1" and "Đầu mối: IT; Liên hệ: máy lẻ 200" in s.text for s in doc.sections)


def test_unsupported_extension(tmp_path):
    path = tmp_path / "x.txt"
    path.write_text("abc")
    with pytest.raises(ValueError):
        load_document(path)
