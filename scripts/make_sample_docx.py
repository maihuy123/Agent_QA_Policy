"""Tạo file DOCX mẫu (hướng dẫn onboarding) để kiểm thử loader DOCX.

Chạy: python scripts/make_sample_docx.py
"""

from datetime import datetime
from pathlib import Path

import docx

OUTPUT = Path(__file__).resolve().parent.parent / "data" / "raw" / "huong_dan_onboarding.docx"

CONTENT = [
    ("Ngày làm việc đầu tiên", [
        "Nhân viên mới có mặt tại quầy lễ tân tầng 3 lúc 8 giờ 30 sáng ngày đầu tiên và mang theo CCCD bản gốc.",
        "Phòng Nhân sự hướng dẫn ký hợp đồng, chụp ảnh làm thẻ nhân viên và giới thiệu với nhóm làm việc.",
        "Bộ phận IT bàn giao laptop, tài khoản email và hướng dẫn cài đặt VPN trong buổi chiều.",
    ]),
    ("Hồ sơ cần nộp", [
        "Trong vòng 7 ngày làm việc kể từ ngày đầu tiên, nhân viên nộp cho Phòng Nhân sự: sơ yếu lý lịch có xác nhận, "
        "bản sao bằng cấp có công chứng, giấy khám sức khỏe trong vòng 6 tháng, 2 ảnh 3x4 và số tài khoản ngân hàng để nhận lương.",
    ]),
    ("Người hướng dẫn (buddy)", [
        "Mỗi nhân viên mới được phân công một người hướng dẫn trong nhóm trong suốt 2 tháng đầu. "
        "Người hướng dẫn giúp nhân viên mới làm quen công việc, quy trình và trả lời các câu hỏi hằng ngày.",
    ]),
    ("Thời gian thử việc", [
        "Thời gian thử việc là 2 tháng đối với vị trí chuyên viên và 1 tháng đối với vị trí nhân viên hỗ trợ. "
        "Lương thử việc bằng 85% lương chính thức. Cuối kỳ thử việc, quản lý trực tiếp đánh giá và Phòng Nhân sự thông báo kết quả.",
    ]),
    ("Đào tạo bắt buộc", [
        "Trong tháng đầu tiên, nhân viên mới hoàn thành ba khóa học trực tuyến trên hệ thống LMS: "
        "Văn hóa công ty, Nhận thức bảo mật thông tin và An toàn lao động.",
    ]),
]

TABLE = [
    ("Đầu mối", "Liên hệ"),
    ("Phòng Nhân sự", "hr@saomai.example - máy lẻ 102"),
    ("Bộ phận IT", "it-support@saomai.example - máy lẻ 200"),
    ("Lễ tân", "máy lẻ 100"),
]


def main() -> None:
    document = docx.Document()
    document.add_heading("Hướng dẫn onboarding cho nhân viên mới", level=0)
    for heading, paragraphs in CONTENT:
        document.add_heading(heading, level=1)
        for text in paragraphs:
            document.add_paragraph(text)
    document.add_heading("Đầu mối liên hệ", level=1)
    table = document.add_table(rows=0, cols=2)
    for left, right in TABLE:
        row = table.add_row().cells
        row[0].text, row[1].text = left, right

    props = document.core_properties
    props.title = "Hướng dẫn onboarding cho nhân viên mới"
    props.author = "Phòng Nhân sự"
    props.modified = datetime(2026, 4, 1)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"Đã tạo {OUTPUT}")


if __name__ == "__main__":
    main()
